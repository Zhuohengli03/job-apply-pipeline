#!/usr/bin/env python3
"""投递追踪 Excel 工具。

用法：
  python3 tracker_update.py init  <xlsx>
  python3 tracker_update.py apply <xlsx> <ops.json> [--backup-dir DIR] [--tag TAG]
  python3 tracker_update.py stats <xlsx>
  python3 tracker_update.py search <xlsx> <关键词> [关键词...]

ops.json 是一个列表，每项是以下两种之一：
  {"op": "append", "sheet": "投递追踪",
   "row": {"公司": "...", "平台": "...", "职位": "...", "日期": "2026-10-29", "状态": "已提交", ...}}
  {"op": "update", "sheet": "投递追踪",
   "match": {"公司": "T-Mobile", "职位": "Associate Product Manager"},   # 子串匹配，全部满足；匹配到多行会报错
   "set": {"状态": "已提交"},
   "append_note": "10/29：Candidate Home 显示 Review"}                   # 可选，追加到「备注」末尾
列名按表头文字匹配，不写死列号。
"""
import copy
import datetime
import json
import os
import shutil
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HEADERS = ["公司", "平台", "职位", "办公（远程/混合/线下）", "日期", "状态", "网址", "备注",
           "使用简历（文件路径）", "Cover Letter（文件路径）", "投递入口网址 / 登录账号"]
SHEETS = ["投递追踪", "BOSS直聘", "猎聘", "实习僧"]
COUNTED = {"已提交", "已投递", "已拒", "进行中", "面试中", "测评中", "Offer"}


def init(path):
    if os.path.exists(path):
        sys.exit(f"已存在，不覆盖：{path}")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "B2"
        for col, w in zip("ABCDEFGHIJK", [22, 18, 30, 18, 12, 10, 30, 60, 30, 30, 36]):
            ws.column_dimensions[col].width = w
    wb.save(path)
    print("created", path)


def header_map(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def find_col(hdr, name):
    if name in hdr:
        return hdr[name]
    for k, v in hdr.items():  # 宽松匹配，如「办公」→「办公（远程/混合/线下）」
        if k.startswith(name):
            return v
    raise KeyError(f"表头里没有列：{name}（现有：{list(hdr)}）")


def last_row(ws):
    r = ws.max_row
    while r > 1 and all(ws.cell(r, c).value in (None, "") for c in range(1, ws.max_column + 1)):
        r -= 1
    return r


def apply(path, ops_path, backup_dir=None, tag="update"):
    ops = json.load(open(ops_path, encoding="utf-8"))
    if backup_dir:
        os.makedirs(backup_dir, exist_ok=True)
        stamp = datetime.date.today().isoformat()
        b = os.path.join(backup_dir, f"{os.path.splitext(os.path.basename(path))[0]}_backup_{stamp}_{tag}.xlsx")
        shutil.copy(path, b)
        print("backup", b)
    wb = openpyxl.load_workbook(path)
    before = list(wb.sheetnames)
    for op in ops:
        ws = wb[op["sheet"]]
        hdr = header_map(ws)
        if op["op"] == "append":
            src = last_row(ws)
            dst = src + 1
            for c in range(1, ws.max_column + 1):
                s = ws.cell(src, c)
                if s.has_style and src > 1:
                    ws.cell(dst, c)._style = copy.copy(s._style)
            if src > 1 and ws.row_dimensions[src].height:
                ws.row_dimensions[dst].height = ws.row_dimensions[src].height
            for k, v in op["row"].items():
                ws.cell(dst, find_col(hdr, k)).value = v
            print("append", op["sheet"], dst, op["row"].get("公司"), op["row"].get("职位"))
        elif op["op"] == "update":
            hits = []
            for r in range(2, ws.max_row + 1):
                if all(str(want) in str(ws.cell(r, find_col(hdr, k)).value or "") for k, want in op["match"].items()):
                    hits.append(r)
            if len(hits) != 1:
                sys.exit(f"update 匹配到 {len(hits)} 行（需要恰好 1 行）：{op['match']} → {hits}")
            r = hits[0]
            for k, v in op.get("set", {}).items():
                ws.cell(r, find_col(hdr, k)).value = v
            if op.get("append_note"):
                c = ws.cell(r, find_col(hdr, "备注"))
                c.value = (str(c.value) + " ｜" if c.value else "") + op["append_note"]
            print("update", op["sheet"], r, op["match"])
        else:
            sys.exit(f"未知 op：{op['op']}")
    wb.save(path)
    after = openpyxl.load_workbook(path).sheetnames
    assert after == before, f"工作表变了：{before} → {after}"
    print("ok, sheets:", after)


def stats(path):
    wb = openpyxl.load_workbook(path)
    total = 0
    for name in wb.sheetnames:
        ws = wb[name]
        hdr = header_map(ws)
        if "状态" not in hdr or "平台" not in hdr:
            continue
        cnt, ea, rej, active = 0, 0, 0, []
        for r in range(2, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                continue
            s = str(ws.cell(r, hdr["状态"]).value or "")
            if s in COUNTED:
                cnt += 1
                if "LinkedIn Easy" in str(ws.cell(r, hdr["平台"]).value or ""):
                    ea += 1
                if s == "已拒":
                    rej += 1
                if s in ("进行中", "面试中", "测评中", "Offer"):
                    active.append(f"{ws.cell(r, 1).value}｜{ws.cell(r, find_col(hdr, '职位')).value}｜{s}")
        print(f"[{name}] 已投 {cnt}（其中 LinkedIn Easy Apply {ea}，官网/ATS {cnt - ea}），已拒 {rej}")
        for a in active:
            print("   在流程中:", a)
        total += cnt
    print("合计（含 BOSS 等打招呼表，按需扣除）:", total)


def search(path, kws):
    wb = openpyxl.load_workbook(path)
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(min_row=2):
            s = " | ".join(str(c.value) for c in row if c.value is not None)
            if any(k.lower() in s.lower() for k in kws):
                print(f"[{name}] 行{row[0].row}: {s[:300]}")


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a:
        sys.exit(__doc__)
    cmd = a[0]
    if cmd == "init":
        init(a[1])
    elif cmd == "apply":
        bd = a[a.index("--backup-dir") + 1] if "--backup-dir" in a else None
        tg = a[a.index("--tag") + 1] if "--tag" in a else "update"
        apply(a[1], a[2], bd, tg)
    elif cmd == "stats":
        stats(a[1])
    elif cmd == "search":
        search(a[1], a[2:])
    else:
        sys.exit(__doc__)
