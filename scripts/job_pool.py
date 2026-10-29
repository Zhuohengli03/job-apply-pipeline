#!/usr/bin/env python3
"""岗位库：把各来源的岗位收进一张表，去重、记住看过没有，每天只看新增的。

用法：
  python3 job_pool.py init   <岗位库.xlsx>
  python3 job_pool.py ingest-simplify <岗位库.xlsx> [--repo New-Grad-Positions] [--repo Summer2027-Internships]
        [--categories Product,AI/ML/Data] [--title-include REGEX] [--title-exclude REGEX]
        [--locations REGEX] [--days 14] [--listings 本地listings.json] [--tracker 投递追踪.xlsx]
  python3 job_pool.py add    <岗位库.xlsx> <jobs.json> [--tracker 投递追踪.xlsx]
  python3 job_pool.py new    <岗位库.xlsx> [--limit 50]
  python3 job_pool.py mark   <岗位库.xlsx> <岗位ID> <状态> [备注]
  python3 job_pool.py stats  <岗位库.xlsx>

来源说明：
  ingest-simplify 读 GitHub 上 SimplifyJobs 维护的应届/实习岗位清单（每天更新的结构化 JSON，
  含类别、地点、发布日期、担保信息）。写明「U.S. Citizenship is Required」或「Does Not Offer
  Sponsorship」的岗位自动标「不投」并写原因（档案里不需要担保的用户可加 --keep-no-sponsor）。
  add 接收其他来源（LinkedIn、Greenhouse/Ashby/Lever API、Handshake…）整理好的 JSON 列表，每项字段：
    {"id": "来源内唯一ID或链接", "source": "LinkedIn", "company": "...", "title": "...",
     "category": "...", "locations": "...", "posted": "2026-10-29", "sponsorship": "...", "url": "..."}

状态：新 / 候选 / 不投 / 已投 / 限投候选 / 已关闭。--tracker 会按「公司+职位」比对投递追踪表，已投过的直接标「已投」。
"""
import datetime
import difflib
import json
import re
import sys
import urllib.request

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

SHEET = "岗位库"
HEADERS = ["岗位ID", "来源", "公司", "职位", "类别", "地点", "发布日期", "担保信息",
           "申请链接", "首次收录", "状态", "打分", "备注", "投递表对应"]
NO_SPONSOR = {"U.S. Citizenship is Required", "Does Not Offer Sponsorship"}
TODAY = datetime.date.today().isoformat()


def norm(s):
    return re.sub(r"[^0-9a-z一-鿿]", "", str(s or "").lower())


def company_key(s):
    s = re.sub(r"[（(].*?[)）]", "", str(s or ""))  # 去掉括号里的说明
    s = re.sub(r"\b(inc|llc|ltd|corp|corporation|co|company|group|technologies|technology)\b\.?", "", s, flags=re.I)
    return norm(s)


def args_get(a, flag, default=None, multi=False):
    vals = [a[i + 1] for i, x in enumerate(a) if x == flag and i + 1 < len(a)]
    if multi:
        return vals or default
    return vals[-1] if vals else default


def init(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "C2"
    for col, w in zip("ABCDEFGHIJKLMN", [16, 12, 22, 40, 14, 26, 11, 18, 40, 11, 9, 7, 50, 14]):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:N1"
    wb.save(path)
    print("created", path)


def load(path):
    try:
        wb = openpyxl.load_workbook(path)
    except FileNotFoundError:
        init(path)
        wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    hdr = {c.value: c.column for c in ws[1]}
    return wb, ws, hdr


def existing_keys(ws, hdr):
    ids, combos = set(), set()
    for r in range(2, ws.max_row + 1):
        i = ws.cell(r, hdr["岗位ID"]).value
        if not i:
            continue
        ids.add(str(i))
        combos.add((company_key(ws.cell(r, hdr["公司"]).value), norm(ws.cell(r, hdr["职位"]).value)))
        u = ws.cell(r, hdr["申请链接"]).value
        if u:
            ids.add(norm(str(u).split("?")[0]))
    return ids, combos


def tracker_index(path):
    """读投递追踪表所有投递工作表的（公司, 职位, 行号, 表名）。"""
    if not path:
        return []
    wb = openpyxl.load_workbook(path, read_only=True)
    out = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = ws.iter_rows(values_only=True)
        hdr = next(rows, None)
        if not hdr or "公司" not in hdr or "职位" not in hdr:
            continue
        ci, ti = hdr.index("公司"), hdr.index("职位")
        si = hdr.index("状态") if "状态" in hdr else None
        for n, row in enumerate(rows, start=2):
            if row[ci]:
                out.append((company_key(row[ci]), norm(row[ti]), n, name, row[si] if si is not None else ""))
    return out


def match_tracker(idx, company, title):
    ck, tk = company_key(company), norm(title)
    if not ck:
        return None
    for c, t, n, sheet, st in idx:
        if not c or not (ck in c or c in ck):
            continue
        if t and (tk in t or t in tk or difflib.SequenceMatcher(None, tk, t).ratio() >= 0.75):
            return f"{sheet}!{n}（{st}）"
    return None


def append_jobs(path, jobs, tracker=None):
    wb, ws, hdr = load(path)
    ids, combos = existing_keys(ws, hdr)
    idx = tracker_index(tracker)
    added = {"新": 0, "不投": 0, "已投": 0}
    dup = 0
    for j in jobs:
        jid = str(j["id"])
        ukey = norm(str(j.get("url", "")).split("?")[0])
        combo = (company_key(j["company"]), norm(j["title"]))
        if jid in ids or (ukey and ukey in ids) or combo in combos:
            dup += 1
            continue
        status, note = "新", j.get("note", "")
        hit = match_tracker(idx, j["company"], j["title"])
        if hit:
            status, note = "已投", (note + " 投递表已有：" + hit).strip()
        elif j.get("auto_skip"):
            status, note = "不投", j["auto_skip"]
        row = {"岗位ID": jid, "来源": j.get("source", ""), "公司": j["company"], "职位": j["title"],
               "类别": j.get("category", ""), "地点": j.get("locations", ""), "发布日期": j.get("posted", ""),
               "担保信息": j.get("sponsorship", ""), "申请链接": j.get("url", ""), "首次收录": TODAY,
               "状态": status, "备注": note, "投递表对应": hit or ""}
        ws.append([row.get(h, "") for h in HEADERS])
        ids.add(jid)
        if ukey:
            ids.add(ukey)
        combos.add(combo)
        added[status] += 1
    wb.save(path)
    print(f"新增 {sum(added.values())}（新 {added['新']}，自动不投 {added['不投']}，已在投递表 {added['已投']}），重复跳过 {dup}")


def ingest_simplify(a):
    path = a[1]
    repos = args_get(a, "--repo", ["New-Grad-Positions"], multi=True)
    cats = [c.strip() for c in args_get(a, "--categories", "Product").split(",") if c.strip()]
    inc = args_get(a, "--title-include")
    exc = args_get(a, "--title-exclude", r"\b(senior|sr\.?|staff|principal|lead|director|manager ii|iii)\b")
    loc = args_get(a, "--locations")
    days = int(args_get(a, "--days", "14"))
    local = args_get(a, "--listings")
    keep_no_sponsor = "--keep-no-sponsor" in a
    cutoff = datetime.datetime.now().timestamp() - days * 86400
    jobs = []
    for repo in repos:
        if local:
            data = json.load(open(local, encoding="utf-8"))
        else:
            url = f"https://raw.githubusercontent.com/SimplifyJobs/{repo}/dev/.github/scripts/listings.json"
            data = json.load(urllib.request.urlopen(url, timeout=60))
        for x in data:
            if not (x.get("active") and x.get("is_visible", True)):
                continue
            if x.get("date_posted", 0) < cutoff:
                continue
            title = x.get("title", "")
            cat_ok = (not cats) or x.get("category") in cats
            inc_ok = bool(inc and re.search(inc, title, re.I))
            if not (cat_ok or inc_ok):
                continue
            if exc and re.search(exc, title, re.I):
                continue
            locs = "; ".join(x.get("locations") or [])
            if loc and not re.search(loc, locs, re.I):
                continue
            sp = x.get("sponsorship", "")
            jobs.append({
                "id": "simplify:" + x["id"], "source": f"GitHub/Simplify {repo}", "company": x.get("company_name", ""),
                "title": title, "category": x.get("category", ""), "locations": locs,
                "posted": datetime.datetime.fromtimestamp(x["date_posted"], datetime.timezone.utc).date().isoformat(),
                "sponsorship": sp, "url": x.get("url", ""),
                "auto_skip": (f"清单标注：{sp}" if sp in NO_SPONSOR and not keep_no_sponsor else ""),
            })
        if local:
            break
    print(f"筛出 {len(jobs)} 个（{days} 天内，类别 {cats}{'，标题含 ' + inc if inc else ''}）")
    append_jobs(path, jobs, args_get(a, "--tracker"))


def add(a):
    jobs = json.load(open(a[2], encoding="utf-8"))
    append_jobs(a[1], jobs, args_get(a, "--tracker"))


def new(a):
    wb, ws, hdr = load(a[1])
    limit = int(args_get(a, "--limit", "50"))
    rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, hdr["状态"]).value == "新"]
    rows.sort(key=lambda r: str(ws.cell(r, hdr["发布日期"]).value or ""), reverse=True)
    print(f"待看 {len(rows)} 个（显示前 {min(limit, len(rows))}）")
    for r in rows[:limit]:
        g = lambda h: ws.cell(r, hdr[h]).value or ""
        print(f"{g('岗位ID')} | {g('发布日期')} | {g('公司')} | {g('职位')} | {g('地点')[:40]} | {g('担保信息')} | {g('申请链接')}")


def mark(a):
    path, jid, status = a[1], a[2], a[3]
    note = a[4] if len(a) > 4 else ""
    wb, ws, hdr = load(path)
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, hdr["岗位ID"]).value) == jid:
            ws.cell(r, hdr["状态"]).value = status
            if note:
                old = ws.cell(r, hdr["备注"]).value
                ws.cell(r, hdr["备注"]).value = (str(old) + " ｜" if old else "") + f"{TODAY[5:]}：{note}"
            wb.save(path)
            print("ok", jid, status)
            return
    sys.exit(f"找不到 {jid}")


def stats(a):
    wb, ws, hdr = load(a[1])
    import collections
    c = collections.Counter(ws.cell(r, hdr["状态"]).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value)
    s = collections.Counter(ws.cell(r, hdr["来源"]).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value)
    print("按状态:", dict(c))
    print("按来源:", dict(s))


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a:
        sys.exit(__doc__)
    {"init": lambda: init(a[1]), "ingest-simplify": lambda: ingest_simplify(a), "add": lambda: add(a),
     "new": lambda: new(a), "mark": lambda: mark(a), "stats": lambda: stats(a)}.get(a[0], lambda: sys.exit(__doc__))()
